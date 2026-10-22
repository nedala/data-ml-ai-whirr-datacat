"""Live warehouse tests against a throwaway Postgres (skipped unless WBTEST_PG=1):

    docker network create bsp-analyst-wbtest-net
    docker run -d --name bsp-analyst-wbtest-pg --network bsp-analyst-wbtest-net \\
      -e POSTGRES_USER=analyst -e POSTGRES_PASSWORD=… -e POSTGRES_DB=bsp_edw postgres:16-alpine
    docker run --rm --network bsp-analyst-wbtest-net -e WBTEST_PG=1 -e PG_HOST=bsp-analyst-wbtest-pg \\
      -e PG_PASS=… -v $PWD:/repo --entrypoint bash bsp/appliance:latest \\
      -c 'cp -r /repo/build/workbench /tmp/wb && cd /tmp/wb && python -m pytest -q tests/test_warehouse_live.py'
"""
from tests import _env  # noqa: F401  (must be first: sets env)
from tests._env import H, PG_LIVE, TMP, wait_job

import csv
import io
import os
import unittest

from fastapi.testclient import TestClient

LIVE = os.environ.get("WBTEST_PG") == "1"
LAKE = TMP / "lake"


@unittest.skipUnless(LIVE, "set WBTEST_PG=1 and PG_* to run against a throwaway Postgres")
class WarehouseLive(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls._saved = {k: os.environ.get(k) for k in ("PG_HOST", "PG_PORT", "PG_DB", "PG_USER", "PG_PASS")}
        os.environ.update(PG_LIVE)
        for k in cls._saved:
            if k not in PG_LIVE:
                os.environ.pop(k, None)
        import app as wbapp
        from wb import warehouse as W
        cls.W = W
        cls.c = TestClient(wbapp.app)
        conn = W.connect()
        with conn, conn.cursor() as c:
            c.execute("DROP SCHEMA IF EXISTS landing CASCADE")
            c.execute("DROP SCHEMA IF EXISTS gold CASCADE")
            c.execute("CREATE SCHEMA gold")
            c.execute("CREATE TABLE gold.kpi AS SELECT g AS id, g * 1.5 AS premium, DATE '2024-01-01' + g AS d FROM generate_series(1, 2500) g")
        conn.close()
        (LAKE / "landing" / "claims.csv").write_text(
            "Claim ID,Paid,Loss Date,Open,Note\nC1,10.5,2024-01-03,yes,\"a,b\"\nC2,,2024-02-04,no,\nC3,abc,2024-03-05,Y,x\n")

    @classmethod
    def tearDownClass(cls):
        for k, v in cls._saved.items():
            if v is None:
                os.environ.pop(k, None)
            else:
                os.environ[k] = v

    def load(self, path, table, mode="replace", columns=None, options=None):
        body = {"root": "lake", "path": path, "table": table, "mode": mode, "columns": columns or [], "options": options or {}}
        r = self.c.post("/ingress/api/load", json=body, headers=H)
        self.assertEqual(r.status_code, 200, r.text)
        return wait_job(self.c, "/ingress", r.json()["job"]["id"])

    def q(self, sql):
        r = self.c.post("/ingress/api/query", json={"sql": sql}, headers=H)
        self.assertEqual(r.status_code, 200, r.text)
        return r.json()

    def test_01_status(self):
        s = self.c.get("/ingress/api/status").json()
        self.assertTrue(s["warehouse"]["ok"], s)

    def test_02_load_with_type_overrides(self):
        cols = [{"name": "Claim ID", "target": "claim_id", "type": "text"},
                {"name": "Paid", "target": "paid", "type": "number"},
                {"name": "Loss Date", "target": "loss_date", "type": "date"},
                {"name": "Open", "target": "is_open", "type": "boolean"},
                {"name": "Note", "target": "note", "type": "text", "include": False}]
        s = self.load("landing/claims.csv", "claims", columns=cols)
        self.assertEqual(s["status"], "succeeded", s)
        self.assertEqual(s["result"]["rows"], 3)
        log = self.c.get(f"/ingress/api/jobs/{s['id']}/log").json()["lines"]
        self.assertTrue(any("couldn't be read as number" in ln for ln in log), log)
        r = self.q("SELECT claim_id, paid, loss_date, is_open FROM landing.claims ORDER BY 1")
        self.assertEqual(r["rows"], [["C1", 10.5, "2024-01-03", True], ["C2", None, "2024-02-04", False], ["C3", None, "2024-03-05", True]])
        self.assertEqual([c["type"] for c in r["columns"]], ["text", "double", "date", "boolean"])

    def test_03_append_and_same_shape_replace_keeps_views(self):
        s = self.load("landing/claims.csv", "claims2")
        self.assertEqual(s["status"], "succeeded", s)
        s = self.load("landing/claims.csv", "claims2", mode="append")
        self.assertEqual(self.q("SELECT count(*) FROM landing.claims2")["rows"], [[6]])
        conn = self.W.connect()
        with conn, conn.cursor() as c:
            c.execute("CREATE VIEW landing.v_claims2 AS SELECT * FROM landing.claims2")
        conn.close()
        s = self.load("landing/claims.csv", "claims2", mode="replace")
        self.assertEqual(s["status"], "succeeded", s)  # same columns → TRUNCATE, the view survives
        self.assertEqual(self.q("SELECT count(*) FROM landing.v_claims2")["rows"], [[3]])
        s = self.load("landing/claims.csv", "claims2", mode="replace",
                      columns=[{"name": "Claim ID", "target": "claim_id", "type": "text"}])
        self.assertEqual(s["status"], "failed")
        self.assertIn("depend", s["error"])
        s = self.load("landing/claims.csv", "claims", mode="append",
                      columns=[{"name": "Note", "target": "nope", "type": "text"}])
        self.assertEqual(s["status"], "failed")
        self.assertIn("nope", s["error"])

    def test_04_excel_and_parquet(self):
        import openpyxl
        import pyarrow as pa
        import pyarrow.parquet as pq
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["Treaty", "Limit", "Inception"])
        ws.append(["T1", 1000000, "2024-01-01"])
        ws.append(["T2", 2500000.5, None])
        wb.save(LAKE / "landing" / "treaties.xlsx")
        s = self.load("landing/treaties.xlsx", "treaties")
        self.assertEqual((s["status"], s["result"]["rows"]), ("succeeded", 2), s)
        pq.write_table(pa.table({"id": [1, 2], "tags": [["a"], ["b", "c"]]}), LAKE / "landing" / "t.parquet")
        s = self.load("landing/t.parquet", "tagged")
        self.assertEqual(s["status"], "succeeded", s)
        r = self.q("SELECT id, tags FROM landing.tagged ORDER BY id")
        self.assertEqual(r["rows"][1], [2, ["b", "c"]])

    def test_05_read_only_enforced(self):
        for sql in ("SELECT 1; DROP TABLE landing.claims", "DELETE FROM landing.claims"):
            self.assertEqual(self.c.post("/ingress/api/query", json={"sql": sql}, headers=H).status_code, 400)
        r = self.c.post("/ingress/api/query", json={"sql": "WITH d AS (DELETE FROM landing.claims RETURNING *) SELECT * FROM d"}, headers=H)
        self.assertEqual(r.status_code, 400)
        self.assertIn("read-only", r.json()["detail"])
        r = self.c.post("/ingress/api/query", json={"sql": "SELECT pg_sleep(0.01), count(*) FROM landing.claims"}, headers=H)
        self.assertEqual(r.status_code, 200)
        self.assertEqual(self.q("SELECT count(*) FROM landing.claims")["rows"], [[3]])

    def test_06_query_limit(self):
        r = self.q("SELECT * FROM gold.kpi")
        self.assertEqual(len(r["rows"]), 200)
        self.assertTrue(r["truncated"])

    def test_07_exports(self):
        for fmt in ("csv", "excel", "parquet"):
            r = self.c.post("/ingress/api/export", json={"schema": "gold", "table": "kpi", "format": fmt}, headers=H)
            self.assertEqual(r.status_code, 200, r.text)
            s = wait_job(self.c, "/ingress", r.json()["job"]["id"])
            self.assertEqual(s["status"], "succeeded", s)
            self.assertEqual(s["result"]["rows"], 2500)
            d = self.c.get(f"/ingress/api/export/{s['result']['token']}")
            self.assertEqual(d.status_code, 200)
            if fmt == "csv":
                rows = list(csv.reader(io.StringIO(d.text)))
                self.assertEqual(rows[0], ["id", "premium", "d"])
                self.assertEqual(rows[1], ["1", "1.5", "2024-01-02"])
            elif fmt == "excel":
                import openpyxl
                ws = openpyxl.load_workbook(io.BytesIO(d.content), read_only=True).active
                self.assertEqual(ws.max_row, 2501)
            else:
                import pyarrow.parquet as pq
                t = pq.read_table(io.BytesIO(d.content))
                self.assertEqual((t.num_rows, t.column_names), (2500, ["id", "premium", "d"]))
        body = {"sql": "SELECT * FROM gold.kpi WHERE id < 10", "format": "csv", "destination": "lake", "filename": "small kpi"}
        r = self.c.post("/ingress/api/export", json=body, headers=H)
        s = wait_job(self.c, "/ingress", r.json()["job"]["id"])
        self.assertEqual(s["result"]["path"], "exports/small kpi.csv")
        self.assertTrue((LAKE / "exports" / "small kpi.csv").exists())
        self.assertEqual(self.c.post("/ingress/api/export", json=body, headers=H).status_code, 409)
        self.assertEqual(self.c.post("/ingress/api/export", json={**body, "overwrite": True}, headers=H).status_code, 200)

    def test_08_catalogue_and_drop(self):
        sch = {s["schema"]: s for s in self.c.get("/ingress/api/warehouse/schemas").json()["schemas"]}
        self.assertIn("landing", sch)
        self.assertGreaterEqual(sch["landing"]["tables"], 1)
        t = self.c.get("/ingress/api/warehouse/tables", params={"schema": "gold"}).json()["tables"]
        self.assertEqual([(x["name"], x["kind"]) for x in t], [("kpi", "table")])
        d = self.c.get("/ingress/api/warehouse/table", params={"schema": "gold", "table": "kpi", "rows": 5}).json()
        self.assertEqual((d["row_count"], len(d["rows"]), d["columns"][0]["name"]), (2500, 5, "id"))
        self.assertTrue(self.c.get("/ingress/api/table-exists", params={"table": "treaties"}).json()["exists"])
        r = self.c.delete("/ingress/api/warehouse/landing/treaties", headers=H)
        self.assertEqual(r.status_code, 200, r.text)
        self.assertFalse(self.c.get("/ingress/api/table-exists", params={"table": "treaties"}).json()["exists"])
        self.assertEqual(self.c.delete("/ingress/api/warehouse/landing/claims2", headers=H).status_code, 400)  # view depends on it


if __name__ == "__main__":
    unittest.main()
