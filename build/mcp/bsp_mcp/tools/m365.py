"""Microsoft 365 via Microsoft Graph — Outlook mail, calendar, Teams, OneDrive/SharePoint files, people.

Mock-or-live like every connector: bsp_connectors.m365 covers mail + files; calendar, Teams, people
and mail search use direct Graph calls in live mode (same Entra app token) and
bsp_mcp/fixtures/m365_extra.yaml in mock mode.

Read-only by default. Write tools (create draft, post to a Teams channel, upload a file) always
return a PREVIEW when dry_run=true (the default) and refuse to commit unless MCP_M365_WRITE=1.
Nothing ever sends mail.
"""
from __future__ import annotations

import datetime as dt
import html
import io
import json
import re
import uuid
from functools import lru_cache
from pathlib import Path
from typing import Annotated, Any

import yaml
from pydantic import Field

from .. import delegated
from ..common import UserError, clamp, connector_mode, connectors, data_dir, env, md_table, result, tool

G = "m365"
FIXTURES = Path(__file__).resolve().parent.parent / "fixtures" / "m365_extra.yaml"
MAX_TEXT = 40_000


def _mode() -> str:
    return connector_mode("m365")


def _live() -> bool:
    return _mode() == "live"


def write_enabled() -> bool:
    return env("MCP_M365_WRITE") == "1"


def _graph(method: str, path: str, **kw: Any) -> Any:
    from bsp_connectors.m365 import live  # noqa: PLC0415

    return live._req(method, path, **kw)


def _paged(path: str, params: dict | None = None, limit: int = 50) -> list[dict]:
    from bsp_connectors._http import paged  # noqa: PLC0415
    from bsp_connectors.m365 import live  # noqa: PLC0415

    return list(paged(live.base() + path, token=live._token, params=params, limit=limit, system="m365"))


@lru_cache(maxsize=1)
def _fx() -> dict:
    return yaml.safe_load(FIXTURES.read_text()) if FIXTURES.exists() else {}


def mailboxes() -> list[str]:
    configured = [m.strip().lower() for m in env("MCP_MAILBOXES").split(",") if m.strip()]
    if configured:
        return configured
    base = [env("DISPUTES_MAILBOX", "billing-disputes@example.net").lower()]
    return base + (["network-ops@example.net"] if not _live() else [])


def _me(mailbox: str | None = None) -> str | None:
    """Delegated Graph token when the request is about the signed-in user's OWN data
    (mailbox omitted / "me" / their address) and they connected Microsoft 365 at SSO."""
    act = delegated.active()
    return act[1] if act and delegated.wants_self(mailbox, act[0]) else None


MSG_FIELDS = "id,receivedDateTime,from,subject,bodyPreview,hasAttachments,isRead,conversationId"


def _check_mailbox(mb: str | None) -> str:
    mb = (mb or mailboxes()[0]).lower()
    if mb not in mailboxes():
        raise UserError(f"mailbox {mb} is not in the MCP allow-list ({', '.join(mailboxes())}); set MCP_MAILBOXES")
    return mb


def _text(body: dict | str | None) -> str:
    if isinstance(body, dict):
        c = body.get("content") or ""
        if (body.get("contentType") or "").lower() == "html":
            c = re.sub(r"<(script|style)[^>]*>.*?</\1>", " ", c, flags=re.S | re.I)
            c = re.sub(r"<br\s*/?>|</p>|</div>|</tr>", "\n", c, flags=re.I)
            c = html.unescape(re.sub(r"<[^>]+>", " ", c))
        return re.sub(r"[ \t]+", " ", c).strip()
    return str(body or "")


def _msg_brief(m: dict) -> dict:
    f = (m.get("from") or {})
    f = f.get("emailAddress", f)
    return {"id": m.get("id"), "received": m.get("receivedDateTime"), "from": f.get("address"), "subject": m.get("subject"),
            "preview": (m.get("bodyPreview") or "")[:160], "hasAttachments": m.get("hasAttachments"), "isRead": m.get("isRead")}


# ── mail ─────────────────────────────────────────────────────────────────────
@tool(G, "List Outlook messages", open_world=True)
def m365_mail_list(
    mailbox: Annotated[str | None, Field(description="Mailbox address (allow-listed; default the disputes mailbox)")] = None,
    folder: Annotated[str, Field(description="Folder name, e.g. inbox")] = "inbox",
    since: Annotated[str | None, Field(description="ISO date/time; only messages received after")] = None,
    unread_only: Annotated[bool, Field(description="Only unread messages")] = False,
    top: Annotated[int, Field(description="Max messages (1–100)")] = 25,
) -> dict:
    """Newest-first messages in YOUR mailbox (delegated, when you signed in with Microsoft and leave mailbox
    empty or "me") or in an allow-listed shared mailbox. Read-only: does not mark anything read."""
    if tok := _me(mailbox):
        flt = " and ".join(([f"receivedDateTime ge {since}"] if since else []) + (["isRead eq false"] if unread_only else []))
        n = clamp(top, 1, 100)
        params = {"$top": str(n), "$orderby": "receivedDateTime desc", "$select": MSG_FIELDS, **({"$filter": flt} if flt else {})}
        items = delegated.graph(tok, "GET", f"/me/mailFolders/{folder}/messages", params, limit=n)
        rows = [_msg_brief(m) for m in items]
        return result(f"{len(rows)} message(s) in your {folder}\n\n" + md_table(rows, 25), mode="delegated", messages=rows)
    mb = _check_mailbox(mailbox)
    rows = [_msg_brief(m) for m in connectors("m365").list_messages(mb, folder=folder, since=since, unread_only=unread_only,
                                                                    top=clamp(top, 1, 100))]
    return result(f"{len(rows)} message(s) in {mb}/{folder}\n\n" + md_table(rows, 25), mode=_mode(), messages=rows)


@tool(G, "Search Outlook messages", open_world=True)
def m365_mail_search(
    query: Annotated[str, Field(description="Words to find in subject/body/sender, e.g. 'late fee 3100008373'")],
    mailbox: Annotated[str | None, Field(description="Mailbox address (allow-listed)")] = None,
    top: Annotated[int, Field(description="Max results (1–50)")] = 20,
) -> dict:
    """Full-text search of YOUR mailbox (delegated; mailbox empty or "me") or an allow-listed shared mailbox."""
    n = clamp(top, 1, 50)
    if tok := _me(mailbox):
        items = delegated.graph(tok, "GET", "/me/messages", {"$search": f'"{query}"', "$top": str(n), "$select": MSG_FIELDS}, limit=n)
        rows = [_msg_brief(m) for m in items]
        return result(f"{len(rows)} match(es) for {query!r} in your mailbox\n\n" + md_table(rows), mode="delegated", messages=rows)
    mb = _check_mailbox(mailbox)
    if _live():
        items = _paged(f"/users/{mb}/messages", {"$search": f'"{query}"', "$top": str(n)}, n)
    else:
        words = [w.lower() for w in query.split() if w]
        items = []
        for m in connectors("m365").list_messages(mb, folder="inbox", top=500):
            hay = " ".join([m.get("subject") or "", _text(m.get("body")), json.dumps(m.get("from") or {})]).lower()
            if all(w in hay for w in words):
                items.append(m)
        items = items[:n]
    rows = [_msg_brief(m) for m in items]
    return result(f"{len(rows)} match(es) for {query!r} in {mb}\n\n" + md_table(rows), mode=_mode(), messages=rows)


@tool(G, "Read one Outlook message", open_world=True)
def m365_mail_read(
    message_id: Annotated[str, Field(description="Message id (from m365_mail_list / m365_mail_search)")],
    mailbox: Annotated[str | None, Field(description="Mailbox address (allow-listed)")] = None,
) -> dict:
    """Subject, sender, recipients, plain-text body and attachment list of one message."""
    if tok := _me(mailbox):
        m = delegated.graph(tok, "GET", f"/me/messages/{message_id}", {"$expand": "attachments($select=id,name,contentType,size)"})
        mode = "delegated"
    else:
        mb = _check_mailbox(mailbox)
        m = connectors("m365").get_message(mb, message_id)
        mode = _mode()
    body = _text(m.get("body"))[:MAX_TEXT]
    atts = [{"id": a.get("id"), "name": a.get("name"), "contentType": a.get("contentType"), "size": a.get("size")}
            for a in m.get("attachments") or []]
    return result(f"**{m.get('subject')}** from {_msg_brief(m)['from']} at {m.get('receivedDateTime')}\n\n{body[:1500]}",
                  mode=mode, message={**_msg_brief(m), "to": m.get("toRecipients"), "body_text": body,
                                         "attachments": atts, "conversationId": m.get("conversationId")})


@tool(G, "Create an Outlook draft (preview by default; never sends)", read_only=False, open_world=True)
def m365_mail_create_draft(
    to: Annotated[list[str], Field(description="Recipient addresses")],
    subject: Annotated[str, Field(description="Subject")],
    body: Annotated[str, Field(description="Plain-text body")],
    mailbox: Annotated[str | None, Field(description="Mailbox to create the draft in (allow-listed)")] = None,
    dry_run: Annotated[bool, Field(description="true = preview only (default); false = create the draft (needs MCP_M365_WRITE=1)")] = True,
) -> dict:
    """Prepare a reply/notification as a DRAFT in Outlook for a human to review and send. Nothing is sent."""
    mb = _check_mailbox(mailbox)
    draft = {"subject": subject, "body": {"contentType": "Text", "content": body},
             "toRecipients": [{"emailAddress": {"address": a}} for a in to]}
    if dry_run:
        return result(f"PREVIEW (not created) — draft in {mb} to {', '.join(to)}: {subject}", mode=_mode(), preview=draft,
                      write_enabled=write_enabled())
    if not write_enabled():
        raise UserError("M365 writes are disabled (set MCP_M365_WRITE=1 on the MCP server to allow creating drafts)")
    if _live():
        d = _graph("POST", f"/users/{mb}/messages", json=draft)
        return result(f"draft created in {mb}: {subject}", mode=_mode(), draft_id=d.get("id"), webLink=d.get("webLink"))
    p = data_dir() / "m365_mock" / "drafts.jsonl"
    p.parent.mkdir(parents=True, exist_ok=True)
    did = "draft-" + uuid.uuid4().hex[:10]
    with p.open("a") as f:
        f.write(json.dumps({"id": did, "mailbox": mb, **draft, "created": dt.datetime.utcnow().isoformat()}) + "\n")
    return result(f"draft created in {mb} (mock): {subject}", mode=_mode(), draft_id=did)


# ── calendar ─────────────────────────────────────────────────────────────────
def _mock_events(user: str, start: dt.datetime, end: dt.datetime) -> list[dict]:
    today = dt.datetime.now(dt.timezone.utc).replace(hour=0, minute=0, second=0, microsecond=0)
    out = []
    for e in (_fx().get("calendars") or {}).get(user.lower(), []):
        d = today + dt.timedelta(days=int(e.get("day", 0)))
        s = d.replace(hour=int(e["start"][:2]), minute=int(e["start"][3:]))
        t = d.replace(hour=int(e["end"][:2]), minute=int(e["end"][3:]))
        if t > start and s < end:
            out.append({"subject": e["subject"], "start": s.isoformat(), "end": t.isoformat(), "location": e.get("location"),
                        "organizer": e.get("organizer")})
    return sorted(out, key=lambda x: x["start"])


def _window(start: str | None, days: int) -> tuple[dt.datetime, dt.datetime]:
    s = dt.datetime.fromisoformat(start.replace("Z", "+00:00")) if start else dt.datetime.now(dt.timezone.utc)
    if s.tzinfo is None:
        s = s.replace(tzinfo=dt.timezone.utc)
    return s, s + dt.timedelta(days=clamp(days, 1, 31))


def _events(user: str, s: dt.datetime, e: dt.datetime) -> list[dict]:
    if tok := _me(user):
        items = delegated.graph(tok, "GET", "/me/calendarView", {"startDateTime": s.isoformat(), "endDateTime": e.isoformat(),
                                                                 "$orderby": "start/dateTime", "$top": "100"}, limit=200)
        return [{"subject": i.get("subject"), "start": (i.get("start") or {}).get("dateTime"), "end": (i.get("end") or {}).get("dateTime"),
                 "location": (i.get("location") or {}).get("displayName"),
                 "organizer": ((i.get("organizer") or {}).get("emailAddress") or {}).get("address")} for i in items]
    if not _live():
        return _mock_events(user, s, e)
    items = _paged(f"/users/{user}/calendarView", {"startDateTime": s.isoformat(), "endDateTime": e.isoformat(),
                                                    "$orderby": "start/dateTime", "$top": "100"}, 200)
    return [{"subject": i.get("subject"), "start": (i.get("start") or {}).get("dateTime"), "end": (i.get("end") or {}).get("dateTime"),
             "location": (i.get("location") or {}).get("displayName"),
             "organizer": ((i.get("organizer") or {}).get("emailAddress") or {}).get("address")} for i in items]


@tool(G, "List calendar events", open_world=True)
def m365_calendar_list_events(
    user: Annotated[str, Field(description="User's e-mail address; omit or 'me' for your own calendar")] = "me",
    start: Annotated[str | None, Field(description="ISO start (default now)")] = None,
    days: Annotated[int, Field(description="Window length in days (1–31)")] = 7,
) -> dict:
    """A user's calendar events in a time window (Graph calendarView; needs Calendars.Read in live mode)."""
    s, e = _window(start, days)
    rows = _events(user, s, e)
    return result(f"{len(rows)} event(s) for {user}\n\n" + md_table(rows, 30), mode="delegated" if _me(user) else _mode(), events=rows)


@tool(G, "Who am I on Microsoft 365", open_world=True)
def m365_whoami() -> dict:
    """Whether the M365 tools act as you (delegated, after "Sign in with Microsoft" with Graph consent) and as whom."""
    act = delegated.active()
    if not act:
        who = delegated.caller()
        why = ("you have not connected Microsoft 365 — sign out and use 'Sign in with Microsoft' (consent to the "
               "requested permissions)") if who else "this client is not a signed-in harness session"
        return result(f"Not acting as a user: {why}. Shared, allow-listed mailboxes still work ({_mode()} mode).",
                      mode=_mode(), delegated=False)
    me = delegated.graph(act[1], "GET", "/me", {"$select": "displayName,mail,userPrincipalName,jobTitle,department"})
    return result(f"Acting as **{me.get('displayName')}** ({me.get('userPrincipalName')})", mode="delegated", delegated=True, me=me)


@tool(G, "List my Teams chats", open_world=True)
def m365_teams_my_chats(top: Annotated[int, Field(description="Max chats (1–50)")] = 15) -> dict:
    """Your recent Teams chats with their last message preview (delegated; needs Chat.Read)."""
    tok = _me()
    if not tok:
        raise UserError("needs your Microsoft 365 connection — sign in with Microsoft (Chat.Read) and try again")
    n = clamp(top, 1, 50)
    items = delegated.graph(tok, "GET", "/me/chats", {"$top": str(n), "$expand": "lastMessagePreview", "$orderby": "lastMessagePreview/createdDateTime desc"}, limit=n)
    rows = [{"topic": c.get("topic") or c.get("chatType"), "type": c.get("chatType"),
             "last": ((c.get("lastMessagePreview") or {}).get("body") or {}).get("content", "")[:120],
             "at": (c.get("lastMessagePreview") or {}).get("createdDateTime")} for c in items]
    return result(f"{len(rows)} chat(s)\n\n" + md_table(rows, 20), mode="delegated", chats=rows)


@tool(G, "Find free time across calendars", open_world=True)
def m365_calendar_find_free_time(
    users: Annotated[list[str], Field(description="Attendee e-mail addresses")],
    duration_minutes: Annotated[int, Field(description="Meeting length")] = 30,
    start: Annotated[str | None, Field(description="ISO start of the search window (default now)")] = None,
    days: Annotated[int, Field(description="Days to search (1–14)")] = 3,
    work_hours_utc: Annotated[str, Field(description="Working hours HH-HH in UTC")] = "13-21",
) -> dict:
    """Slots where every attendee is free (working hours only), computed from their calendars."""
    s, e = _window(start, min(days, 14))
    busy = []
    for u in users:
        for ev in _events(u, s, e):
            busy.append((dt.datetime.fromisoformat(str(ev["start"]).replace("Z", "+00:00")[:25]),
                         dt.datetime.fromisoformat(str(ev["end"]).replace("Z", "+00:00")[:25])))
    busy = [(a if a.tzinfo else a.replace(tzinfo=dt.timezone.utc), b if b.tzinfo else b.replace(tzinfo=dt.timezone.utc)) for a, b in busy]
    h0, h1 = (int(x) for x in work_hours_utc.split("-"))
    step, dur = dt.timedelta(minutes=30), dt.timedelta(minutes=clamp(duration_minutes, 15, 480))
    t = s.replace(minute=0 if s.minute < 30 else 30, second=0, microsecond=0) + step
    slots = []
    while t + dur <= e and len(slots) < 10:
        if h0 <= t.hour and (t + dur).hour <= h1 and t.weekday() < 5 and not any(a < t + dur and b > t for a, b in busy):
            slots.append({"start": t.isoformat(), "end": (t + dur).isoformat()})
        t += step
    return result(f"{len(slots)} free slot(s) for {', '.join(users)}\n\n" + md_table(slots), mode=_mode(), slots=slots)


# ── Teams ────────────────────────────────────────────────────────────────────
@tool(G, "List Teams and channels", open_world=True)
def m365_teams_list(team_id: Annotated[str | None, Field(description="Team id to list its channels; omit to list teams")] = None) -> dict:
    """Teams (or one team's channels) visible to the app (Team.ReadBasic.All / Channel.ReadBasic.All live)."""
    if team_id:
        if _live():
            rows = [{"id": c["id"], "displayName": c.get("displayName"), "description": c.get("description")}
                    for c in _paged(f"/teams/{team_id}/channels")]
        else:
            t = next((t for t in _fx().get("teams") or [] if t["id"] == team_id), None)
            if t is None:
                raise UserError(f"team {team_id} not found")
            rows = [{"id": c["id"], "displayName": c["displayName"]} for c in t.get("channels") or []]
        return result(f"{len(rows)} channel(s)\n\n" + md_table(rows), mode=_mode(), channels=rows)
    if _live():
        rows = [{"id": g["id"], "displayName": g.get("displayName"), "description": g.get("description")}
                for g in _paged("/teams", {"$top": "100"}, 200)]
    else:
        rows = [{"id": t["id"], "displayName": t["displayName"], "description": t.get("description")} for t in _fx().get("teams") or []]
    return result(f"{len(rows)} team(s)\n\n" + md_table(rows), mode=_mode(), teams=rows)


@tool(G, "Read Teams channel messages", open_world=True)
def m365_teams_read_channel(
    team_id: Annotated[str, Field(description="Team id")],
    channel_id: Annotated[str, Field(description="Channel id")],
    top: Annotated[int, Field(description="Max messages (1–50)")] = 20,
) -> dict:
    """Recent messages in a channel (ChannelMessage.Read.All live)."""
    n = clamp(top, 1, 50)
    if _live():
        items = _paged(f"/teams/{team_id}/channels/{channel_id}/messages", {"$top": str(n)}, n)
        rows = [{"from": ((i.get("from") or {}).get("user") or {}).get("displayName"), "at": i.get("createdDateTime"),
                 "text": _text(i.get("body"))[:1000]} for i in items]
    else:
        t = next((t for t in _fx().get("teams") or [] if t["id"] == team_id), {})
        ch = next((c for c in t.get("channels") or [] if c["id"] == channel_id), None)
        if ch is None:
            raise UserError(f"channel {channel_id} not found in {team_id}")
        today = dt.datetime.now(dt.timezone.utc).date()
        rows = []
        for m in ch.get("messages") or []:
            day, hm = m["at"].split()
            rows.append({"from": m["from"], "at": f"{today + dt.timedelta(days=int(day))}T{hm}:00Z", "text": m["text"]})
        rows = rows[-n:]
    return result(f"{len(rows)} message(s)\n\n" + "\n".join(f"- {r['at']} **{r['from']}**: {r['text']}" for r in rows),
                  mode=_mode(), messages=rows)


@tool(G, "Post to a Teams channel (preview by default)", read_only=False, open_world=True)
def m365_teams_post_message(
    team_id: Annotated[str, Field(description="Team id")],
    channel_id: Annotated[str, Field(description="Channel id")],
    text: Annotated[str, Field(description="Message text")],
    dry_run: Annotated[bool, Field(description="true = preview only (default); false = post (needs MCP_M365_WRITE=1)")] = True,
) -> dict:
    """Post a message to a Teams channel. Disabled unless MCP_M365_WRITE=1; previews by default."""
    if dry_run:
        return result(f"PREVIEW (not posted) to {team_id}/{channel_id}: {text[:200]}", mode=_mode(),
                      preview={"team_id": team_id, "channel_id": channel_id, "text": text}, write_enabled=write_enabled())
    if not write_enabled():
        raise UserError("Teams posting is disabled (set MCP_M365_WRITE=1 on the MCP server)")
    if _live():
        d = _graph("POST", f"/teams/{team_id}/channels/{channel_id}/messages", json={"body": {"content": text}})
        return result("posted", mode=_mode(), message_id=d.get("id"))
    p = data_dir() / "m365_mock" / "teams_posts.jsonl"
    p.parent.mkdir(parents=True, exist_ok=True)
    with p.open("a") as f:
        f.write(json.dumps({"team_id": team_id, "channel_id": channel_id, "text": text,
                            "at": dt.datetime.utcnow().isoformat()}) + "\n")
    return result("posted (mock)", mode=_mode())


# ── files (OneDrive / SharePoint) ────────────────────────────────────────────
def _file_brief(f: dict) -> dict:
    return {"id": f.get("id"), "name": f.get("name"), "path": f.get("parentPath"), "size": f.get("size"),
            "modified": f.get("lastModifiedDateTime"), "webUrl": f.get("webUrl")}


@tool(G, "List SharePoint / OneDrive files", open_world=True)
def m365_files_list(
    folder: Annotated[str, Field(description="Folder path in the document library, e.g. 'CCAI/CallLogs' ('' = root)")] = "",
    site: Annotated[str | None, Field(description="SharePoint site (default SHAREPOINT_SITE)")] = None,
    recursive: Annotated[bool, Field(description="Include sub-folders")] = False,
) -> dict:
    """Files in a document-library folder."""
    rows = [_file_brief(f) for f in connectors("m365").list_files(site, folder, recursive=recursive)][:500]
    return result(f"{len(rows)} file(s) in /{folder}\n\n" + md_table(rows, 30), mode=_mode(), files=rows)


@tool(G, "Search SharePoint / OneDrive files", open_world=True)
def m365_files_search(
    query: Annotated[str, Field(description="Words in the file name (live: Graph drive search incl. content)")],
    site: Annotated[str | None, Field(description="SharePoint site (default SHAREPOINT_SITE)")] = None,
    top: Annotated[int, Field(description="Max results (1–100)")] = 25,
) -> dict:
    """Find files by name (and by content in live/delegated mode). Without a site and with your Microsoft 365
    connection, searches YOUR OneDrive and the SharePoint files shared with you."""
    n = clamp(top, 1, 100)
    if not site and (tok := _me()):
        q = query.replace("'", "''")
        items = delegated.graph(tok, "GET", f"/me/drive/root/search(q='{q}')", {"$top": str(n)}, limit=n)
        rows = [{"id": i.get("id"), "name": i.get("name"), "path": ((i.get("parentReference") or {}).get("path") or ""),
                 "size": i.get("size"), "modified": i.get("lastModifiedDateTime"), "webUrl": i.get("webUrl")} for i in items]
        return result(f"{len(rows)} file(s) in your OneDrive match {query!r}\n\n" + md_table(rows), mode="delegated", files=rows)
    if _live():
        from bsp_connectors.m365 import live  # noqa: PLC0415

        q = query.replace("'", "''")
        items = _paged(f"/sites/{live.site_id(site)}/drive/root/search(q='{q}')", {"$top": str(n)}, n)
        rows = [{"id": i.get("id"), "name": i.get("name"), "path": ((i.get("parentReference") or {}).get("path") or ""),
                 "size": i.get("size"), "modified": i.get("lastModifiedDateTime"), "webUrl": i.get("webUrl")} for i in items]
    else:
        words = [w.lower() for w in query.split()]
        rows = [_file_brief(f) for f in connectors("m365").list_files(site, "", recursive=True)
                if all(w in ((f.get("parentPath") or "") + "/" + f.get("name", "")).lower() for w in words)][:n]
    return result(f"{len(rows)} file(s) match {query!r}\n\n" + md_table(rows), mode=_mode(), files=rows)


def extract_text(name: str, data: bytes) -> str:
    low = name.lower()
    if low.endswith(".pdf"):
        from pypdf import PdfReader  # noqa: PLC0415

        return "\n\n".join((p.extract_text() or "") for p in PdfReader(io.BytesIO(data)).pages)
    if low.endswith(".docx"):
        import docx  # noqa: PLC0415

        d = docx.Document(io.BytesIO(data))
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            parts += [" | ".join(c.text for c in r.cells) for r in t.rows]
        return "\n".join(parts)
    if low.endswith((".xlsx", ".xlsm")):
        import openpyxl  # noqa: PLC0415

        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        out = []
        for ws in wb.worksheets:
            out.append(f"## sheet {ws.title}")
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i >= 500:
                    out.append("…")
                    break
                out.append(",".join("" if v is None else str(v) for v in row))
        return "\n".join(out)
    if low.endswith(".pptx"):
        import zipfile  # noqa: PLC0415

        z = zipfile.ZipFile(io.BytesIO(data))
        slides = sorted((n for n in z.namelist() if re.match(r"ppt/slides/slide\d+\.xml$", n)),
                        key=lambda n: int(re.findall(r"\d+", n)[-1]))
        return "\n\n".join(f"## slide {i + 1}\n" + " ".join(re.findall(r"<a:t>([^<]*)</a:t>", z.read(s).decode("utf-8", "ignore")))
                           for i, s in enumerate(slides))
    if low.endswith((".txt", ".csv", ".md", ".json", ".vtt", ".srt", ".log", ".yaml", ".yml", ".xml", ".html", ".eml")):
        return data.decode("utf-8", errors="replace")
    raise UserError(f"unsupported file type for text extraction: {name} (docx, xlsx, pptx, pdf, txt/csv/md/json/vtt)")


@tool(G, "Read the text of a SharePoint / OneDrive file", open_world=True)
def m365_files_read(
    item_id: Annotated[str, Field(description="File id (from m365_files_list / m365_files_search)")],
    name: Annotated[str | None, Field(description="File name (helps pick the parser; looked up when omitted)")] = None,
    site: Annotated[str | None, Field(description="SharePoint site (default SHAREPOINT_SITE)")] = None,
    max_chars: Annotated[int, Field(description="Truncate the text (1k–40k)")] = 20000,
) -> dict:
    """Download one file and return its text (docx, xlsx, pptx, pdf, txt/csv/md/json/vtt)."""
    m = connectors("m365")
    if not name:
        hit = next((f for f in m.list_files(site, "", recursive=True) if f.get("id") == item_id), None)
        name = (hit or {}).get("name") or "file.txt"
    data = m.download_file(site, item_id)
    text = extract_text(name, data)
    cap = clamp(max_chars, 1000, MAX_TEXT)
    return result(f"{name}: {len(text):,} chars{' (truncated)' if len(text) > cap else ''}\n\n{text[:1200]}", mode=_mode(),
                  name=name, text=text[:cap], truncated=len(text) > cap)


@tool(G, "Upload a text file to SharePoint (preview by default)", read_only=False, open_world=True)
def m365_files_upload(
    folder: Annotated[str, Field(description="Target folder path")],
    name: Annotated[str, Field(description="File name, e.g. 'summary.md'")],
    content: Annotated[str, Field(description="Text content")],
    site: Annotated[str | None, Field(description="SharePoint site (default SHAREPOINT_SITE)")] = None,
    dry_run: Annotated[bool, Field(description="true = preview only (default); false = upload (needs MCP_M365_WRITE=1)")] = True,
) -> dict:
    """Upload a small text file. Disabled unless MCP_M365_WRITE=1; previews by default."""
    if dry_run:
        return result(f"PREVIEW (not uploaded): /{folder}/{name} ({len(content.encode())} bytes)", mode=_mode(),
                      preview={"folder": folder, "name": name, "bytes": len(content.encode()), "head": content[:500]},
                      write_enabled=write_enabled())
    if not write_enabled():
        raise UserError("SharePoint uploads are disabled (set MCP_M365_WRITE=1 on the MCP server)")
    item = connectors("m365").upload_file(site, folder, name, content.encode())
    return result(f"uploaded /{folder}/{name}", mode=_mode(), file=_file_brief(item))


# ── people ───────────────────────────────────────────────────────────────────
@tool(G, "Look up people in the directory", open_world=True)
def m365_people_lookup(query: Annotated[str, Field(description="Name, e-mail, title or department")],
                       top: Annotated[int, Field(description="Max results (1–25)")] = 10) -> dict:
    """Directory search (User.Read.All live): display name, mail, job title, department, office."""
    n = clamp(top, 1, 25)
    if _live():
        q = query.replace('"', "")
        data = _graph("GET", "/users", params={"$search": f'"displayName:{q}" OR "mail:{q}"', "$top": str(n),
                                                "$select": "id,displayName,mail,jobTitle,department,officeLocation"},
                      headers={"ConsistencyLevel": "eventual"}) or {}
        rows = data.get("value") or []
    else:
        ql = query.lower()
        rows = [p for p in _fx().get("people") or [] if any(ql in str(v).lower() for v in p.values())][:n]
    return result(f"{len(rows)} person(s)\n\n" + md_table(rows), mode=_mode(), people=rows)
