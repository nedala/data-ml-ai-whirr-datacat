"""In-process MCP client against the server (mock connectors), bearer auth over HTTP, audit log, rag group."""
import asyncio
import json
import sys
import types

from starlette.testclient import TestClient


def call(server, tname, **args):
    from fastmcp import Client

    async def go():
        async with Client(server) as c:
            return await c.call_tool(tname, args, raise_on_error=False)
    r = asyncio.run(go())
    text = "\n".join(x.text for x in r.content if getattr(x, "type", "") == "text")
    return r.is_error, (json.loads(text) if text.startswith("{") else text)


def test_m365_mock_tools_and_write_gate(monkeypatch):
    from bsp_mcp.server import build
    srv = build()
    err, d = call(srv, "m365_mail_list", top=3)
    assert not err and d["mode"] == "mock" and len(d["messages"]) == 3
    err, d = call(srv, "m365_mail_search", query="refund")
    assert not err and d["messages"]
    mid = d["messages"][0]["id"]
    err, d = call(srv, "m365_mail_read", message_id=mid)
    assert not err and d["message"]["body_text"]
    err, d = call(srv, "m365_mail_list", mailbox="ceo@example.net")
    assert err and "allow-list" in d
    err, d = call(srv, "m365_mail_create_draft", to=["x@example.net"], subject="s", body="b")
    assert not err and d["summary"].startswith("PREVIEW")
    err, d = call(srv, "m365_mail_create_draft", to=["x@example.net"], subject="s", body="b", dry_run=False)
    assert err and "MCP_M365_WRITE" in d
    monkeypatch.setenv("MCP_M365_WRITE", "1")
    err, d = call(srv, "m365_mail_create_draft", to=["x@example.net"], subject="s", body="b", dry_run=False)
    assert not err and d["draft_id"].startswith("draft-")
    err, d = call(srv, "m365_teams_list")
    assert not err and d["teams"]
    err, d = call(srv, "m365_teams_read_channel", team_id="team-billing", channel_id="ch-billing-general")
    assert not err and d["messages"]
    err, d = call(srv, "m365_calendar_find_free_time", users=["avery.brooks@example.net"], days=5)
    assert not err and "slots" in d
    err, d = call(srv, "m365_people_lookup", query="network")
    assert not err and d["people"][0]["displayName"] == "Jordan Patel"
    err, d = call(srv, "m365_files_search", query="ccai")
    assert not err and d["files"]
    f = d["files"][0]
    err, d = call(srv, "m365_files_read", item_id=f["id"], name=f["name"])
    assert not err and d["text"]


def test_atlan_powerbi_mock():
    from bsp_mcp.server import build
    srv = build()
    err, d = call(srv, "atlan_glossary_lookup", term="churn")
    assert not err and d["terms"]
    err, d = call(srv, "powerbi_list_workspaces")
    assert not err and d["workspaces"]


def test_file_text_extraction():
    import io

    import docx
    import openpyxl
    from bsp_mcp.tools.m365 import extract_text
    d = docx.Document(); d.add_paragraph("Fiber outage credit policy"); b = io.BytesIO(); d.save(b)
    assert "outage credit" in extract_text("p.docx", b.getvalue())
    wb = openpyxl.Workbook(); wb.active.append(["ban", "amt"]); wb.active.append(["3100008373", 14.99]); b = io.BytesIO(); wb.save(b)
    assert "3100008373,14.99" in extract_text("x.xlsx", b.getvalue())


def test_http_bearer_auth_and_audit(tmp_path, monkeypatch):
    from bsp_mcp import audit
    from bsp_mcp.server import create_app
    app = create_app(token="t0k")
    with TestClient(app) as c:
        assert c.get("/healthz").json()["status"] == "ok"
        assert c.get("/mcp/info").status_code == 200
        body = {"jsonrpc": "2.0", "id": 1, "method": "tools/list", "params": {}}
        hdr = {"Accept": "application/json, text/event-stream", "Content-Type": "application/json"}
        assert c.post("/mcp", json=body, headers=hdr).status_code == 401
        assert c.post("/mcp", json=body, headers={**hdr, "Authorization": "Bearer nope"}).status_code == 401
        init = {"jsonrpc": "2.0", "id": 0, "method": "initialize", "params": {"protocolVersion": "2025-06-18", "capabilities": {},
                                                                               "clientInfo": {"name": "t", "version": "1"}}}
        r = c.post("/mcp/", json=init, headers={**hdr, "Authorization": "Bearer t0k"})
        assert r.status_code == 200 and r.json()["result"]["serverInfo"]["name"] == "dataplatform"
        r = c.post("/mcp/m365", json=body, headers={**hdr, "Authorization": "Bearer t0k"})
        names = [t["name"] for t in r.json()["result"]["tools"]]
        assert names and all(n.startswith("m365_") for n in names)
        call_body = {"jsonrpc": "2.0", "id": 2, "method": "tools/call", "params": {"name": "m365_people_lookup", "arguments": {"query": "lee"}}}
        r = c.post("/mcp", json=call_body, headers={**hdr, "Authorization": "Bearer t0k"})
        assert r.status_code == 200 and not r.json()["result"]["isError"]
        assert c.post("/mcp/nope", json=body, headers=hdr).status_code == 404
    rows = audit.recent(10)
    assert rows and rows[0]["tool"] == "m365_people_lookup" and rows[0]["status"] == "ok"


def test_rag_group_wraps_bsp_rag_tools(monkeypatch):
    from bsp_mcp.common import REGISTRY
    from bsp_mcp.tools import rag
    mod = types.ModuleType("bsp_rag.tools")

    def rag_search(query: str, collection: str | None = None, k: int = 5) -> dict:
        """Search the knowledge base."""
        return {"results": [{"text": "Reset the ONT by…", "source": "runbook.md", "q": query}]}

    def rag_list_collections() -> list:
        """Collections."""
        return ["runbooks"]
    mod.rag_search, mod.rag_list_collections = rag_search, rag_list_collections
    pkg = types.ModuleType("bsp_rag"); pkg.tools = mod
    monkeypatch.setitem(sys.modules, "bsp_rag", pkg)
    monkeypatch.setitem(sys.modules, "bsp_rag.tools", mod)
    monkeypatch.setattr(rag, "_done", None)
    before = list(REGISTRY.tools)
    REGISTRY.tools[:] = [t for t in REGISTRY.tools if t.group != "rag"]
    try:
        assert rag.register() == ["rag_search", "rag_list_collections"]
        from bsp_mcp.server import build
        err, d = call(build(), "rag_search", query="ont reset")
        assert not err and d["results"][0]["source"] == "runbook.md"
    finally:
        REGISTRY.tools[:] = before
        rag._done = None


def test_rag_absent_is_skipped(monkeypatch):
    from bsp_mcp.tools import rag
    monkeypatch.setitem(sys.modules, "bsp_rag.tools", None)
    monkeypatch.setattr(rag, "_done", None)
    assert rag.register() == []
