"""Generate the binary documents of the demo corpus (SYNTHETIC examples):

  corpus/binary/ds-xgs-pon-migration-plan.pdf      (design,   DS-NET-007)
  corpus/binary/rb-storm-restoration-playbook.docx (runbook,  RB-NOC-015)
  corpus/binary/rb-wire-center-contacts.xlsx       (runbook,  RB-FLD-014)
  corpus/binary/trn-front-office-onboarding.pptx   (reference, TRN-FO-001)

Each gets a sidecar `<file>.meta.yaml` with the chunk metadata. The PDF is
written by a tiny built-in PDF writer (no reportlab needed); DOCX/PPTX/XLSX
use python-docx / python-pptx / openpyxl (in build/requirements/rag.txt).

    python build/rag/scripts/make_binaries.py [--out build/rag/corpus/binary]
"""

from __future__ import annotations

import argparse
import textwrap
from pathlib import Path

import yaml

BANNER = "SYNTHETIC EXAMPLE - illustrative content for the DataPlatform appliance demo, not an actual DataPlatform document."

META = {
    "ds-xgs-pon-migration-plan.pdf": {"doc_id": "DS-NET-007", "title": "GPON to XGS-PON Migration Plan", "doc_type": "design",
                                      "system": "Access network - OLT / ONT", "owner": "Network Architecture", "version": "1.1",
                                      "effective_date": "2026-06-15", "acl_group": "engineering",
                                      "tags": ["XGS-PON", "GPON", "migration", "coexistence", "OLT"]},
    "rb-storm-restoration-playbook.docx": {"doc_id": "RB-NOC-015", "title": "Storm Restoration Playbook", "doc_type": "runbook",
                                           "system": "NOC / Field Operations", "owner": "Network Operations Center", "version": "2.0",
                                           "effective_date": "2026-05-01", "acl_group": "noc",
                                           "tags": ["storm", "restoration", "generator", "battery", "power", "outage"]},
    "rb-wire-center-contacts.xlsx": {"doc_id": "RB-FLD-014", "title": "Wire Center Field Contacts", "doc_type": "runbook",
                                     "system": "Field Operations", "owner": "Field Operations", "version": "2026.09",
                                     "effective_date": "2026-09-01", "acl_group": "noc",
                                     "tags": ["wire center", "contacts", "field ops supervisor", "escalation", "OLT"]},
    "trn-front-office-onboarding.pptx": {"doc_id": "TRN-FO-001", "title": "Front-office Agent Onboarding", "doc_type": "reference",
                                         "system": "Care training", "owner": "Customer Care Training", "version": "3",
                                         "effective_date": "2026-04-01", "acl_group": "frontoffice",
                                         "tags": ["training", "onboarding", "billing call", "verification"]},
}


# ── minimal PDF writer ───────────────────────────────────────────────────
def _esc(s: str) -> str:
    return s.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def write_pdf(path: Path, title: str, lines: list[tuple[str, str]], author: str) -> None:
    """lines: (style, text) with style in h1|h2|p|li. Helvetica, US Letter."""
    pages: list[list[str]] = [[]]
    y = 740
    for style, text in lines:
        size, lead, font, indent, width = {"h1": (16, 24, "F2", 0, 70), "h2": (12, 20, "F2", 0, 80),
                                          "p": (10, 13, "F1", 0, 95), "li": (10, 13, "F1", 14, 90)}[style]
        wrapped = textwrap.wrap(text, width) or [""]
        if style in ("h1", "h2"):
            y -= 6
        for i, w in enumerate(wrapped):
            if y < 60:
                pages.append([])
                y = 740
            x = 60 + (indent if (style == "li" and i > 0) else 0)
            pages[-1].append(f"BT /{font} {size} Tf {x} {y} Td ({_esc(w)}) Tj ET")
            y -= lead
        y -= 4
    objs: list[bytes] = []

    def add(b: str | bytes) -> int:
        objs.append(b.encode("latin-1") if isinstance(b, str) else b)
        return len(objs)

    add("<< /Type /Catalog /Pages 2 0 R >>")
    add("PAGES")  # placeholder
    f1 = add("<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>")
    f2 = add("<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica-Bold /Encoding /WinAnsiEncoding >>")
    kids = []
    for ops in pages:
        stream = "\n".join(ops).encode("latin-1", errors="replace")
        c = add(b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream")
        p = add(f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 {f1} 0 R /F2 {f2} 0 R >> >> "
                f"/Contents {c} 0 R >>")
        kids.append(p)
    objs[1] = f"<< /Type /Pages /Kids [{' '.join(f'{k} 0 R' for k in kids)}] /Count {len(kids)} >>".encode()
    info = add(f"<< /Title ({_esc(title)}) /Author ({_esc(author)}) /Subject (SYNTHETIC EXAMPLE) /Producer (bsp_rag make_binaries) >>")
    out = bytearray(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
    offsets = []
    for i, o in enumerate(objs, 1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + o + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R /Info {info} 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    path.write_bytes(bytes(out))


def make_pdf(out: Path) -> None:
    L = [("h1", "GPON to XGS-PON Migration Plan"), ("p", BANNER),
         ("h2", "1 Purpose"),
         ("p", "This plan describes how fiber serving areas migrate from GPON to XGS-PON without new outside plant. "
               "XGS-PON delivers 10 Gbps symmetrical capacity per PON port and enables the Fiber 2 Gig plan (FB-2G)."),
         ("h2", "2 Coexistence design"),
         ("p", "GPON uses 1490 nm downstream and 1310 nm upstream. XGS-PON uses 1577 nm downstream and 1270 nm upstream. "
               "Because the wavelengths differ, both technologies share the same optical distribution network (feeder, "
               "splitters, drops). A coexistence element (CEx) at the OLT combines the GPON and XGS-PON line card ports "
               "onto one feeder fiber."),
         ("p", "The CEx adds about 1.5 dB of insertion loss; serving areas with less than 2 dB of optical budget margin "
               "need a splitter audit before migration."),
         ("h2", "3 Migration steps"),
         ("li", "1. Select PON ports above 70% evening peak utilisation for 14 days, or with Fiber 2 Gig demand."),
         ("li", "2. Install XGS-PON line cards and the CEx in the wire center OLT during a maintenance window."),
         ("li", "3. Patch the feeder through the CEx; existing GPON ONTs keep working on their wavelengths."),
         ("li", "4. Swap customer ONTs to XGS-PON (combo ONTs) only when the customer upgrades to a multi-gig plan or "
                "the ONT fails."),
         ("li", "5. Move remaining GPON customers in waves; retire GPON cards when a port has no GPON ONTs left."),
         ("h2", "4 Customer impact"),
         ("p", "Patching through the CEx causes an outage of under 5 minutes per PON and is done between 00:00 and 05:00. "
               "Customers on the affected PON receive a maintenance notice 72 hours in advance. Maintenance events are "
               "excluded from customer-impacting outage minutes."),
         ("h2", "5 Wave plan 2026"),
         ("p", "Wire centers are identified by CLLI code. Wave 1 (Q3): RLGHNCXA, WLSNNCXA, MNFDOHXA - highest Fiber 2 Gig demand. Wave 2 (Q4): CHVLVAXA, ERIEPAXA, TYLRTXXA. "
               "Wave 3 (2027 H1): remaining wire centers with more than 2,000 fiber subscribers.")]
    write_pdf(out / "ds-xgs-pon-migration-plan.pdf", "GPON to XGS-PON Migration Plan", L, "Network Architecture")


def make_docx(out: Path) -> None:
    import docx

    d = docx.Document()
    d.core_properties.title = "Storm Restoration Playbook"
    d.core_properties.author = "Network Operations Center"
    d.core_properties.subject = "SYNTHETIC EXAMPLE"
    d.core_properties.keywords = "storm, restoration, generator, battery"
    d.add_heading("Storm Restoration Playbook", 0)
    d.add_paragraph(BANNER)
    d.add_heading("Before the storm (T-48 to T-0 hours)", 1)
    for s in ["Confirm fuel levels at every wire center generator are above 80% and schedule top-ups.",
              "Stage portable generators at regional yards for field cabinets with less than 8 hours of battery.",
              "Pre-position splicing crews and spare drop cable in the forecast impact area.",
              "Publish the storm bridge number and the NOC staffing roster; switch Care to the storm IVR script."]:
        d.add_paragraph(s, style="List Number")
    d.add_heading("During restoration", 1)
    for s in ["Prioritise power: deploy a portable generator to any cabinet or hut on battery with less than 2 hours remaining.",
              "Restore central office and OLT power before field plant; an OLT down affects every PON it serves.",
              "Only after power is stable, dispatch splicing crews to fiber cuts, ordered by customers affected.",
              "Do not dispatch individual repair tickets inside the storm polygon until area restoration is confirmed.",
              "Update the ETR on the outage card every 60 minutes; storm ETRs may be given as a range (for example 24-48 hours)."]:
        d.add_paragraph(s, style="List Number")
    d.add_heading("Safety", 1)
    d.add_paragraph("Never work on aerial plant near downed power lines; wait for the electric utility to make the area safe. "
                    "Lightning within 10 miles stops all aerial and bucket-truck work.")
    d.add_heading("Credits and communications", 1)
    d.add_paragraph("Storm outages lasting 24 hours or more receive the automatic prorated outage credit. "
                    "Customers in a FEMA-declared disaster county also receive late-fee reversals for 60 days.")
    t = d.add_table(rows=1, cols=3)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text, t.rows[0].cells[2].text = "Battery remaining", "Action", "Owner"
    for r in [("> 4 hours", "Monitor", "NOC"), ("2-4 hours", "Dispatch generator", "Field Ops supervisor"),
              ("< 2 hours", "Priority generator + notify Regional Manager", "NOC Duty Manager")]:
        cells = t.add_row().cells
        for i, v in enumerate(r):
            cells[i].text = v
    d.save(out / "rb-storm-restoration-playbook.docx")


def make_xlsx(out: Path) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    wb.properties.title = "Wire Center Field Contacts"
    wb.properties.creator = "Field Operations"
    ws = wb.active
    ws.title = "Wire centers"
    ws.append(["Wire center (CLLI)", "City", "Region", "OLTs", "Fiber subscribers", "Copper lines", "Field Ops supervisor", "Supervisor phone",
               "Generator on site", "NOC bridge"])
    rows = [
        ("RLGHNCXA", "Raleigh", "East", 4, 6400, 5200, "Jordan Ellis", "555-0201", "Yes", "555-0199 x214"),
        ("WLSNNCXA", "Wilson", "East", 2, 2900, 3100, "Avery Morgan", "555-0202", "Yes", "555-0199 x233"),
        ("CHVLVAXA", "Charlottesville", "East", 3, 4100, 2800, "Riley Chen", "555-0203", "Yes", "555-0199 x115"),
        ("ERIEPAXA", "Erie", "East", 2, 2300, 3900, "Casey Patel", "555-0204", "No", "555-0199 x302"),
        ("MNFDOHXA", "Mansfield", "Central", 3, 3800, 4400, "Morgan Reyes", "555-0211", "Yes", "555-0199 x381"),
        ("LIMAOHXA", "Lima", "Central", 2, 1900, 3600, "Taylor Brooks", "555-0212", "Yes", "555-0199 x405"),
        ("JPLNMOXA", "Joplin", "Central", 2, 2100, 3300, "Jamie Foster", "555-0213", "No", "555-0199 x268"),
        ("TYLRTXXA", "Tyler", "South", 3, 3300, 2700, "Drew Alvarez", "555-0221", "Yes", "555-0199 x457"),
        ("FTMYFLXA", "Fort Myers", "South", 4, 5200, 2500, "Quinn Harper", "555-0222", "Yes", "555-0199 x122"),
        ("DTHNALXA", "Dothan", "South", 1, 1200, 2900, "Sydney Ward", "555-0223", "No", "555-0199 x310"),
    ]
    for r in rows:
        ws.append(list(r))
    ws2 = wb.create_sheet("Notes")
    ws2.append(["Note"])
    ws2.append([BANNER])
    ws2.append(["Call the Field Ops supervisor for dispatch decisions; use the NOC bridge for SEV1/SEV2 coordination."])
    wb.save(out / "rb-wire-center-contacts.xlsx")


def make_pptx(out: Path) -> None:
    from pptx import Presentation

    prs = Presentation()
    prs.core_properties.title = "Front-office Agent Onboarding"
    prs.core_properties.author = "Customer Care Training"
    slides = [
        ("Front-office Agent Onboarding", [BANNER, "Week 1: billing, technical and retention basics"],
         "Welcome new agents and explain how the knowledge base and agent assist help on every call."),
        ("Every billing call: four things", ["1. Verify the caller (account PIN or last four of the payment method)",
                                             "2. Open the invoice and identify the disputed line",
                                             "3. Check policy in the knowledge base before promising any credit",
                                             "4. Note the account with the outcome and the policy you applied"],
         "Verification comes first because of CPNI rules. Never promise credits beyond your authority limit."),
        ("Technical calls", ["Always check for an outage in the customer's wire center first",
                             "Fiber: read the ONT state and light levels before rebooting",
                             "Only dispatch after remote troubleshooting is documented"], ""),
        ("Retention calls", ["Find the reason, fix the root cause, then present one matching offer",
                             "Never refuse or obstruct a cancellation"], ""),
        ("Using agent assist", ["Suggestions come from runbooks with citations - open the source before acting",
                                "Drafted replies must be reviewed and edited before sending"], ""),
    ]
    for title, bullets, notes in slides:
        s = prs.slides.add_slide(prs.slide_layouts[1])
        s.shapes.title.text = title
        tf = s.placeholders[1].text_frame
        tf.text = bullets[0]
        for b in bullets[1:]:
            tf.add_paragraph().text = b
        if notes:
            s.notes_slide.notes_text_frame.text = notes
    prs.save(out / "trn-front-office-onboarding.pptx")


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default=str(Path(__file__).resolve().parents[1] / "corpus" / "binary"))
    out = Path(ap.parse_args().out)
    out.mkdir(parents=True, exist_ok=True)
    make_pdf(out)
    make_docx(out)
    make_xlsx(out)
    make_pptx(out)
    for name, meta in META.items():
        (out / (name + ".meta.yaml")).write_text("# SYNTHETIC EXAMPLE — sidecar metadata for the ingester\n" +
                                                yaml.safe_dump(meta, sort_keys=False), encoding="utf-8")
    print("wrote", ", ".join(sorted(p.name for p in out.iterdir())))


if __name__ == "__main__":
    main()
